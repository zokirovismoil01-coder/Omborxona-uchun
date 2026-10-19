#!/usr/bin/env python3
"""Baliqxona oylik hisob-kitobi uchun Excel ishchi kitobini yaratadi.

Ishlatish:
    python3 oylik/tools/excel_shablon.py [chiqish.xlsx]

Kitob bir oy uchun mo'ljallangan va quyidagi varaqlardan iborat:
    Xulosa       — oy nomi, yaxlitlash qadami, umumiy natijalar va tekshiruvlar
    Oylik hisob  — har bir ishchining hisob-kitobi (to'lov qaydnomasi)
    Amallar      — avans, mukofot, jarima va to'lovlar jurnali
    Guruhlar     — guruh oyliklari, me'yorlar va tejam tahlili
    Qo'llanma    — foydalanish bo'yicha yo'riqnoma

Hisob qoidasi ilovadagi js/calc.js bilan bir xil (tests/excel_mosligi.py ikkalasini solishtiradi).
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from copy import copy
from datetime import datetime
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

# ---------- Boshlang'ich ma'lumotlar (asl "Baliqxona_oylik.xlsx" faylidan) ----------

OY_NOMI = "Oktabr 2026"
QADAM = 1000          # oylik shu summagacha pastga yaxlitlanadi
ISH_KUNI = 26         # guruhning oylik ish kuni me'yori (o'zgartirish mumkin)

GURUHLAR = [
    # nomi,       asosiy oylik, me'yoriy soni, oyiga ish kuni
    ("Ofitsiant", 4_000_000, 5, ISH_KUNI),
    ("Moyka", 4_000_000, 2, ISH_KUNI),
    ("Shef", 4_000_000, 1, ISH_KUNI),
    ("Uborsa", 3_500_000, 1, ISH_KUNI),
]

ISHCHILAR = [
    # ism, guruh, kelmagan kun, o'tgan oydan qoldiq, izoh
    ("Ofitsiant 1", "Ofitsiant", 0, 0, ""),
    ("Ofitsiant 2", "Ofitsiant", 0, 0, ""),
    ("Ofitsiant 3", "Ofitsiant", 0, 0, ""),
    ("Ofitsiant 5", "Ofitsiant", 0, 0, ""),
    ("Moyka 1", "Moyka", 0, 0, ""),
    ("Moyka 2", "Moyka", 0, 0, ""),
    ("Shef", "Shef", 0, 0, ""),
    ("Uborsa", "Uborsa", 0, 0, ""),
    ("Ali", "Ofitsiant", 0, 0, ""),
    ("Kamola", "Moyka", 0, 0, ""),
]

# Amallar jurnali: (sana, ism, turi, summa, izoh). Bo'sh bo'lsa bitta bo'sh qator qoladi.
AMALLAR = []

AMAL_TURLARI = ["Avans", "Mukofot", "Jarima", "To'lov"]

# ---------- Ko'rinish ----------

FONT = "Arial"
KOK = "0000FF"           # qo'lda kiritiladigan kataklar (asl fayldagidek)
SARLAVHA_FON = "1F4E78"  # asl fayldagi sarlavha rangi
JAMI_FON = "DDEBF7"      # asl fayldagi JAMI qatori rangi
PUL = '#,##0;[Red]\\-#,##0;"-"'
SON = '0;[Red]\\-0;"-"'
SANA = "dd.mm.yyyy"

f_oddiy = Font(name=FONT, size=10)
f_kirit = Font(name=FONT, size=10, color=KOK)
f_qalin = Font(name=FONT, size=10, bold=True)
f_sarlavha = Font(name=FONT, size=14, bold=True, color=SARLAVHA_FON)
f_izoh = Font(name=FONT, size=9, italic=True, color="595959")
f_bolim = Font(name=FONT, size=10, bold=True, color="FFFFFF")
fill_bolim = PatternFill("solid", fgColor=SARLAVHA_FON)
fill_jami = PatternFill("solid", fgColor=JAMI_FON)
ingichka = Side(style="thin", color="9BC2E6")
chegara_jami = Border(top=ingichka, bottom=ingichka)
markaz = Alignment(horizontal="center", vertical="center", wrap_text=True)

MARKAZDA = {"№", "Holat", "Tekshiruv", "Kelmagan kun", "Ish kuni", "Ishlagan kun", "Turi",
            "Me'yoriy soni", "Oyiga ish kuni", "Hozirgi soni"}

ZAXIRA_QATOR = 300  # tekshiruv va formatlash shu qatorgacha qo'llanadi (yangi qatorlar uchun)

# Jadval nomlari
G, I, A = "Guruhlar", "Ishchilar", "Amallar"

# Guruhlar jadvali ustunlari
G_NOM, G_ASOSIY, G_MEYOR, G_KUN = "Guruh", "Asosiy oylik", "Me'yoriy soni", "Oyiga ish kuni"
G_FOND, G_SONI, G_HOLAT, G_KISHI = "Ajratilgan fond", "Hozirgi soni", "Holat", "1 kishiga oylik"
G_TOLIQ, G_DAVOMAT = "To'liq oylik jami", "Davomat bo'yicha"
G_BOSH, G_KELMAGAN = "Bo'sh o'rin tejami", "Kelmagan kunlar tejami"
G_YAXLIT, G_TEJAM = "Yaxlitlash qoldig'i", "Jami tejaldi"

# Ishchilar jadvali ustunlari
I_NO, I_ISM, I_GURUH, I_KELMAGAN = "№", "Ism", "Guruh", "Kelmagan kun"
I_KUN, I_ISHLAGAN, I_KISHI, I_ASOSIY = "Ish kuni", "Ishlagan kun", "1 kishiga oylik", "Asosiy"
I_MUKOFOT, I_JARIMA, I_OTGAN, I_JAMI = "Mukofot", "Jarima", "O'tgan oydan", "Jami hisoblangan"
I_AVANS, I_QOLGA, I_TOLANDI, I_QOLDIQ = "Avans", "Qo'lga beriladi", "To'landi", "Qoldiq"
I_HOLAT, I_IZOH, I_IMZO = "Holat", "Izoh", "Imzo"

# Amallar jadvali ustunlari
A_SANA, A_ISM, A_TURI, A_SUMMA, A_IZOH, A_TEK = "Sana", "Ism", "Turi", "Summa", "Izoh", "Tekshiruv"


def esc(nom):
    """Jadval ustuni nomidagi ' belgisi formulada '' bo'lib yoziladi (Excel qoidasi)."""
    return nom.replace("'", "''")


def ust(jadval, nom):
    """Butun ustunga havola: Ishchilar[Asosiy]"""
    return f"{jadval}[{esc(nom)}]"


def joriy(jadval, nom):
    """Joriy qatordagi katak: Ishchilar[[#This Row],[Asosiy]]"""
    return f"{jadval}[[#This Row],[{esc(nom)}]]"


def ustun_harfi(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def jadval_qur(ws, nom, boshi_qator, ustunlar, qatorlar, uslub="TableStyleMedium2"):
    """ustunlar: [(sarlavha, kenglik, turi, formula|None, format|None, izoh|None)];
    turi — 'kirit' (ko'k, qo'lda) yoki 'formula'. qatorlar — faqat 'kirit' ustunlari qiymatlari."""
    for j, (sarlavha, kenglik, _t, _f, _fmt, izoh) in enumerate(ustunlar):
        c = ws.cell(boshi_qator, j + 1, sarlavha)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        # Jadval uslubini ko'rsatmaydigan dasturlarda (LibreOffice, ba'zi telefon ilovalari)
        # ham sarlavha ko'rinib tursin
        c.fill = fill_bolim
        c.alignment = markaz
        if izoh:
            c.comment = Comment(izoh, "Baliqxona", width=280, height=120)
        ws.column_dimensions[ustun_harfi(j)].width = kenglik
    if not qatorlar:
        qatorlar = [None]  # Excel jadvalida kamida bitta qator bo'lishi kerak
    for i, qiymatlar in enumerate(qatorlar):
        r = boshi_qator + 1 + i
        k = 0
        for j, (_s, _k, turi, formula, fmt, _iz) in enumerate(ustunlar):
            c = ws.cell(r, j + 1)
            if turi == "formula":
                c.value = "=" + formula
                c.font = f_oddiy
            else:
                if qiymatlar is not None:
                    v = qiymatlar[k]
                    c.value = None if v == "" else v
                k += 1
                c.font = f_kirit
            if fmt:
                c.number_format = fmt
            if ustunlar[j][0] in MARKAZDA:
                c.alignment = Alignment(horizontal="center")
    oxiri = boshi_qator + len(qatorlar)
    t = Table(displayName=nom, ref=f"A{boshi_qator}:{ustun_harfi(len(ustunlar) - 1)}{oxiri}")
    t.tableStyleInfo = TableStyleInfo(name=uslub, showRowStripes=True)
    t._initialise_columns()
    # Hisoblanadigan ustun: yangi qo'shilgan qatorlarga formula o'zi tushadi
    for col, (_s, _k, turi, formula, _f, _iz) in zip(t.tableColumns, ustunlar):
        if turi == "formula":
            col.calculatedColumnFormula = TableFormula(attr_text=formula)
    ws.add_table(t)
    ws.row_dimensions[boshi_qator].height = 30
    return oxiri


def zaxira_format(ws, oxiri, ustunlar):
    """Jadval ostidagi bo'sh qatorlarga ham ko'k shrift va son formatini beradi —
    jadval kengayganda yangi qatorlar shu ko'rinishni oladi."""
    for r in range(oxiri + 1, ZAXIRA_QATOR + 1):
        for j, (_s, _k, turi, _f, fmt, _iz) in enumerate(ustunlar):
            c = ws.cell(r, j + 1)
            c.font = f_oddiy if turi == "formula" else f_kirit
            if fmt:
                c.number_format = fmt
            if ustunlar[j][0] in MARKAZDA:
                c.alignment = Alignment(horizontal="center")


def jami_qatori(ws, qator, ustunlar, jadval, yigindi_ustunlar, yorliq_ustun=0):
    """Jadval sarlavhasi ustidagi JAMI qatori. SUBTOTAL filtrlangan qatorlarni hisobga oladi."""
    for j in range(len(ustunlar)):
        c = ws.cell(qator, j + 1)
        c.fill = fill_jami
        c.font = f_qalin
        c.border = chegara_jami
    ws.cell(qator, yorliq_ustun + 1, "JAMI")
    nomlar = [u[0] for u in ustunlar]
    for nom in yigindi_ustunlar:
        j = nomlar.index(nom)
        c = ws.cell(qator, j + 1, f"=SUBTOTAL(109,{ust(jadval, nom)})")
        c.number_format = ustunlar[j][4] or PUL


def shart_rangi(ws, diapazon, formula, fon, shrift):
    ws.conditional_formatting.add(
        diapazon,
        FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=fon, bgColor=fon),
                    font=Font(color=shrift)),
    )


def tekshiruv(ws, diapazon, **kw):
    dv = DataValidation(showErrorMessage=True, allow_blank=True, **kw)
    dv.add(diapazon)
    ws.add_data_validation(dv)


def guruhlar_varagi(gs, guruhlar):
    g = lambda n: joriy(G, n)  # noqa: E731
    ustunlar = [
        (G_NOM, 16, "kirit", None, None,
         "Guruh nomi. Ishchilar shu nom bo'yicha guruhga bog'lanadi."),
        (G_ASOSIY, 14, "kirit", None, PUL, "Bir kishining to'liq oyligi (so'm)."),
        (G_MEYOR, 10, "kirit", None, SON,
         "Guruhda nechta odam ishlashi kerak. Fond = asosiy oylik × me'yoriy soni."),
        (G_KUN, 10, "kirit", None, SON,
         "Bir oydagi ish kunlari me'yori. Kelmagan kunlar shu songa nisbatan ushlab qolinadi."),
        (G_FOND, 14, "formula", f"{g(G_ASOSIY)}*{g(G_MEYOR)}", PUL, None),
        (G_SONI, 10, "formula", f"COUNTIF({ust(I, I_GURUH)},{g(G_NOM)})", SON,
         "«Oylik hisob» varag'idagi shu guruh ishchilari soni."),
        (G_HOLAT, 18, "formula",
         f"IF({g(G_NOM)}=\"\",\"\","
         f"IF(COUNTIF({ust(G, G_NOM)},{g(G_NOM)})>1,\"⚠ Nom takrorlangan\","
         f"IF(N({g(G_KUN)})<=0,\"⚠ Ish kuni yo'q\","
         f"IF({g(G_SONI)}=0,\"Hech kim yo'q\","
         f"IF({g(G_SONI)}<{g(G_MEYOR)},\"Bo'sh o'rin: \"&({g(G_MEYOR)}-{g(G_SONI)}),"
         f"IF({g(G_SONI)}>{g(G_MEYOR)},\"Me'yordan ortiq: \"&({g(G_SONI)}-{g(G_MEYOR)}),"
         f"\"Me'yorda\"))))))",
         None, None),
        (G_KISHI, 14, "formula",
         f"IF({g(G_SONI)}=0,0,MIN({g(G_ASOSIY)},INT({g(G_FOND)}/{g(G_SONI)}/Qadam+1E-9)*Qadam))",
         PUL, "Fond hozirgi odamlarga teng bo'linadi, lekin asosiy oylikdan oshmaydi. "
              "Natija yaxlitlash qadamigacha pastga yaxlitlanadi."),
        (G_TOLIQ, 14, "formula", f"{g(G_KISHI)}*{g(G_SONI)}", PUL,
         "Hamma to'liq ishlaganda guruhga to'lanadigan summa."),
        (G_DAVOMAT, 14, "formula",
         f"SUMIFS({ust(I, I_ASOSIY)},{ust(I, I_GURUH)},{g(G_NOM)})", PUL,
         "Kelmagan kunlar hisobga olingandan keyingi haqiqiy asosiy oylik jami."),
        (G_BOSH, 13, "formula", f"MAX(0,{g(G_MEYOR)}-{g(G_SONI)})*{g(G_ASOSIY)}", PUL,
         "Me'yordan kam odam ishlagani uchun ortib qolgan pul."),
        (G_KELMAGAN, 13, "formula", f"{g(G_TOLIQ)}-{g(G_DAVOMAT)}", PUL,
         "Ishchilar kelmagan kunlar uchun ushlab qolingan pul."),
        (G_YAXLIT, 13, "formula", f"{g(G_FOND)}-{g(G_TOLIQ)}-{g(G_BOSH)}", PUL,
         "Fondni teng bo'lib, pastga yaxlitlaganda qolgan mayda qoldiq."),
        (G_TEJAM, 14, "formula", f"{g(G_FOND)}-{g(G_DAVOMAT)}", PUL, None),
    ]
    gs["A1"] = "Guruhlar: oylik fondi va tejam"
    gs["A1"].font = f_sarlavha
    gs["A2"] = ("Ko'k kataklar qo'lda kiritiladi. Yangi guruh: jadval ostidagi birinchi bo'sh "
                "qatorga nom yozing — jadval o'zi kengayadi.")
    gs["A2"].font = f_izoh
    oxiri = jadval_qur(gs, G, 5, ustunlar, [list(x) for x in guruhlar])
    jami_qatori(gs, 4, ustunlar, G, [G_MEYOR, G_FOND, G_SONI, G_TOLIQ, G_DAVOMAT, G_BOSH,
                                     G_KELMAGAN, G_YAXLIT, G_TEJAM])
    zaxira_format(gs, oxiri, ustunlar)
    gs.freeze_panes = "B6"
    nomlar = [u[0] for u in ustunlar]
    h = ustun_harfi(nomlar.index(G_HOLAT))
    shart_rangi(gs, f"{h}6:{h}{ZAXIRA_QATOR}", f'LEFT(${h}6,1)="⚠"', "FFC7CE", "9C0006")
    shart_rangi(gs, f"{h}6:{h}{ZAXIRA_QATOR}", f'LEFT(${h}6,9)="Me\'yordan"', "FCE4D6", "9C3D00")
    shart_rangi(gs, f"{h}6:{h}{ZAXIRA_QATOR}",
                f'OR(LEFT(${h}6,4)="Bo\'s",LEFT(${h}6,4)="Hech")', "FFF2CC", "7F6000")
    shart_rangi(gs, f"A6:A{ZAXIRA_QATOR}",
                f'AND($A6<>"",COUNTIF($A$6:$A${ZAXIRA_QATOR},$A6)>1)', "FFC7CE", "9C0006")
    tekshiruv(gs, f"B6:C{ZAXIRA_QATOR}", type="whole", operator="greaterThanOrEqual",
              formula1="0", errorTitle="Noto'g'ri qiymat", error="Musbat butun son kiriting.")
    tekshiruv(gs, f"D6:D{ZAXIRA_QATOR}", type="whole", operator="between", formula1="1",
              formula2="31", errorTitle="Noto'g'ri qiymat",
              error="Ish kuni 1 dan 31 gacha bo'lishi kerak.")


def oylik_varagi(om, ishchilar):
    w = lambda n: joriy(I, n)  # noqa: E731

    def amal_summa(turi):
        return (f"SUMIFS({ust(A, A_SUMMA)},{ust(A, A_ISM)},{w(I_ISM)},"
                f"{ust(A, A_TURI)},\"{turi}\")")

    def guruhdan(nom):
        return f"IFERROR(INDEX({ust(G, nom)},MATCH({w(I_GURUH)},{ust(G, G_NOM)},0)),0)"

    ustunlar = [
        (I_NO, 5, "formula", f"ROW()-ROW({I}[#Headers])", "0", None),
        (I_ISM, 22, "kirit", None, None,
         "Ishchining ismi va familiyasi. Har bir ism faqat bir marta yozilsin."),
        (I_GURUH, 13, "kirit", None, None, "Ro'yxatdan tanlang («Guruhlar» varag'idagi nomlar)."),
        (I_KELMAGAN, 9, "kirit", None, SON,
         "Shu oyda ishga chiqmagan kunlar soni. Hamma kun ishlagan bo'lsa 0."),
        (I_KUN, 8, "formula", guruhdan(G_KUN), SON,
         "Guruhning oylik ish kuni me'yori («Guruhlar» varag'idan)."),
        (I_ISHLAGAN, 9, "formula", f"MAX(0,{w(I_KUN)}-{w(I_KELMAGAN)})", SON, None),
        (I_KISHI, 13, "formula", guruhdan(G_KISHI), PUL,
         "Guruh bo'yicha bir kishiga to'g'ri keladigan to'liq oylik."),
        (I_ASOSIY, 13, "formula",
         f"IF({w(I_KISHI)}=0,0,IF({w(I_KELMAGAN)}<=0,{w(I_KISHI)},IF({w(I_KUN)}<=0,0,"
         f"INT({w(I_KISHI)}*{w(I_ISHLAGAN)}/{w(I_KUN)}/Qadam+1E-9)*Qadam)))",
         PUL, "Davomat bo'yicha oylik: 1 kishiga oylik × ishlagan kun ÷ ish kuni, "
              "yaxlitlash qadamigacha pastga yaxlitlanadi."),
        (I_MUKOFOT, 12, "formula", amal_summa("Mukofot"), PUL, "«Amallar» varag'idan avtomatik."),
        (I_JARIMA, 12, "formula", amal_summa("Jarima"), PUL, "«Amallar» varag'idan avtomatik."),
        (I_OTGAN, 12, "kirit", None, PUL,
         "O'tgan oydan qolgan qoldiq. Musbat — ishchiga qarzimiz, manfiy — ishchining qarzi "
         "(avans oyligidan ko'p berilgan). O'tgan oy faylidagi «Qoldiq» ustunidan olinadi."),
        (I_JAMI, 14, "formula",
         f"{w(I_ASOSIY)}+{w(I_MUKOFOT)}-{w(I_JARIMA)}+{w(I_OTGAN)}", PUL, None),
        (I_AVANS, 12, "formula", amal_summa("Avans"), PUL, "«Amallar» varag'idan avtomatik."),
        (I_QOLGA, 14, "formula", f"{w(I_JAMI)}-{w(I_AVANS)}", PUL,
         "Oylik kuni beriladigan summa (avanslar ayirilgan)."),
        (I_TOLANDI, 12, "formula", amal_summa("To'lov"), PUL,
         "«Amallar» varag'idagi «To'lov» yozuvlari."),
        (I_QOLDIQ, 13, "formula", f"{w(I_QOLGA)}-{w(I_TOLANDI)}", PUL,
         "To'lanmay qolgan summa. Manfiy bo'lsa — ishchi qarzdor. Keyingi oyga o'tadi."),
        (I_HOLAT, 17, "formula",
         f"IF({w(I_ISM)}=\"\",\"\","
         f"IF(COUNTIF({ust(G, G_NOM)},{w(I_GURUH)})=0,\"⚠ Guruh topilmadi\","
         f"IF(COUNTIF({ust(I, I_ISM)},{w(I_ISM)})>1,\"⚠ Ism takrorlangan\","
         f"IF(AND({w(I_QOLDIQ)}=0,{w(I_JAMI)}=0,{w(I_AVANS)}=0,{w(I_TOLANDI)}=0),\"—\","
         f"IF({w(I_QOLDIQ)}<0,\"Qarzdor\","
         f"IF({w(I_QOLDIQ)}=0,\"✓ To'landi\","
         f"IF({w(I_TOLANDI)}>0,\"Qisman to'landi\",\"To'lanmagan\")))))))",
         None, None),
        (I_IZOH, 18, "kirit", None, None, None),
        (I_IMZO, 12, "kirit", None, None, "Chop etilgan qaydnomada ishchi imzosi uchun."),
    ]
    qatorlar = [[ism, guruh, kelmagan, otgan, izoh, None]
                for (ism, guruh, kelmagan, otgan, izoh) in ishchilar]

    om["A1"] = '="Oylik hisob — "&OyNomi'
    om["A1"].font = f_sarlavha
    om["A2"] = ("Ko'k kataklar qo'lda kiritiladi. Avans, mukofot, jarima va to'lovlar «Amallar» "
                "varag'iga yoziladi. Yangi ishchi: jadval ostidagi birinchi bo'sh qatorga ism yozing.")
    om["A2"].font = f_izoh
    oxiri = jadval_qur(om, I, 5, ustunlar, qatorlar)
    jami_qatori(om, 4, ustunlar, I, [I_ASOSIY, I_MUKOFOT, I_JARIMA, I_OTGAN, I_JAMI, I_AVANS,
                                     I_QOLGA, I_TOLANDI, I_QOLDIQ], yorliq_ustun=1)
    zaxira_format(om, oxiri, ustunlar)
    om.freeze_panes = "D6"
    nomlar = [u[0] for u in ustunlar]
    h = ustun_harfi(nomlar.index(I_HOLAT))
    rng = f"{h}6:{h}{ZAXIRA_QATOR}"
    shart_rangi(om, rng, f'LEFT(${h}6,1)="⚠"', "FFC7CE", "9C0006")
    shart_rangi(om, rng, f'LEFT(${h}6,1)="✓"', "C6EFCE", "006100")
    shart_rangi(om, rng, f'${h}6="Qarzdor"', "E4DFEC", "5B2C6F")
    shart_rangi(om, rng, f'${h}6="Qisman to\'landi"', "FFEB9C", "9C5700")
    shart_rangi(om, f"B6:B{ZAXIRA_QATOR}",
                f'AND($B6<>"",COUNTIF($B$6:$B${ZAXIRA_QATOR},$B6)>1)', "FFC7CE", "9C0006")
    gh = ustun_harfi(nomlar.index(I_GURUH))
    tekshiruv(om, f"{gh}6:{gh}{ZAXIRA_QATOR}", type="list",
              formula1=f'INDIRECT("{ust(G, G_NOM)}")', errorTitle="Guruh topilmadi",
              error="Guruhni ro'yxatdan tanlang. Yangi guruh kerak bo'lsa, avval «Guruhlar» "
                    "varag'iga qo'shing.")
    kh = ustun_harfi(nomlar.index(I_KELMAGAN))
    tekshiruv(om, f"{kh}6:{kh}{ZAXIRA_QATOR}", type="whole", operator="between", formula1="0",
              formula2="31", errorTitle="Noto'g'ri qiymat",
              error="Kelmagan kunlar 0 dan 31 gacha butun son bo'lishi kerak.")
    om.print_title_rows = "4:5"


def amallar_varagi(am, amallar):
    a = lambda n: joriy(A, n)  # noqa: E731
    turlar = "{" + ",".join(f'"{t}"' for t in AMAL_TURLARI) + "}"
    ustunlar = [
        (A_SANA, 12, "kirit", None, SANA, "Pul berilgan yoki yozilgan sana."),
        (A_ISM, 22, "kirit", None, None, "Ro'yxatdan tanlang («Oylik hisob» varag'idagi ismlar)."),
        (A_TURI, 11, "kirit", None, None,
         "Avans — oy davomida berilgan pul; Mukofot — qo'shiladi; Jarima — ayiriladi; "
         "To'lov — oylik kuni berilgan pul."),
        (A_SUMMA, 13, "kirit", None, PUL, "Har doim musbat son yozing."),
        (A_IZOH, 28, "kirit", None, None, None),
        (A_TEK, 20, "formula",
         f"IF(AND({a(A_ISM)}=\"\",{a(A_SUMMA)}=\"\"),\"\","
         f"IF(COUNTIF({ust(I, I_ISM)},{a(A_ISM)})=0,\"⚠ Ism ro'yxatda yo'q\","
         f"IF(ISNA(MATCH({a(A_TURI)},{turlar},0)),\"⚠ Turi tanlanmagan\","
         f"IF(N({a(A_SUMMA)})<=0,\"⚠ Summa yo'q\",\"✓\"))))",
         None, None),
    ]
    am["A1"] = '="Amallar jurnali — "&OyNomi'
    am["A1"].font = f_sarlavha
    am["A2"] = ("Har bir avans, mukofot, jarima va to'lov alohida qatorga yoziladi. "
                "Summalar «Oylik hisob» varag'iga avtomatik tushadi.")
    am["A2"].font = f_izoh
    oxiri = jadval_qur(am, A, 5, ustunlar, [list(x) for x in amallar], uslub="TableStyleMedium4")
    jami_qatori(am, 4, ustunlar, A, [A_SUMMA], yorliq_ustun=1)
    zaxira_format(am, oxiri, ustunlar)
    am.freeze_panes = "A6"
    shart_rangi(am, f"F6:F{ZAXIRA_QATOR}", 'LEFT($F6,1)="⚠"', "FFC7CE", "9C0006")
    shart_rangi(am, f"C6:C{ZAXIRA_QATOR}", '$C6="Jarima"', "FCE4D6", "9C3D00")
    shart_rangi(am, f"C6:C{ZAXIRA_QATOR}", '$C6="Mukofot"', "E2EFDA", "375623")
    tekshiruv(am, f"B6:B{ZAXIRA_QATOR}", type="list", formula1=f'INDIRECT("{ust(I, I_ISM)}")',
              errorTitle="Ism topilmadi",
              error="Ismni ro'yxatdan tanlang. Yangi ishchini avval «Oylik hisob» varag'iga "
                    "qo'shing.")
    tekshiruv(am, f"C6:C{ZAXIRA_QATOR}", type="list", formula1='"' + ",".join(AMAL_TURLARI) + '"',
              errorTitle="Noto'g'ri tur",
              error="Turini ro'yxatdan tanlang: " + ", ".join(AMAL_TURLARI) + ".")
    tekshiruv(am, f"D6:D{ZAXIRA_QATOR}", type="decimal", operator="greaterThan", formula1="0",
              errorTitle="Noto'g'ri summa", error="Summa musbat son bo'lishi kerak.")
    tekshiruv(am, f"A6:A{ZAXIRA_QATOR}", type="date", operator="greaterThan", formula1="36526",
              errorTitle="Noto'g'ri sana",
              error="Sanani KK.OO.YYYY ko'rinishida kiriting, masalan 15.10.2026.")


def xulosa_varagi(xs, oy_nomi, qadam):
    for col, kenglik in (("A", 34), ("B", 16), ("C", 3), ("D", 34), ("E", 16)):
        xs.column_dimensions[col].width = kenglik
    xs["A1"] = "Baliqxona — oylik xulosa"
    xs["A1"].font = f_sarlavha
    xs["A2"] = ("Faqat ko'k kataklar o'zgartiriladi. Qolgan raqamlar boshqa varaqlardan "
                "avtomatik olinadi.")
    xs["A2"].font = f_izoh

    def yozuv(ref, matn, qalin=False):
        xs[ref] = matn
        xs[ref].font = f_qalin if qalin else f_oddiy

    def qiymat(ref, formula, fmt=PUL, qalin=False):
        xs[ref] = formula
        xs[ref].number_format = fmt
        xs[ref].font = f_qalin if qalin else f_oddiy
        xs[ref].alignment = Alignment(horizontal="right")

    def bolim(ref, matn):
        xs[ref] = matn
        xs[ref].font = f_bolim
        xs[ref].fill = fill_bolim
        keyingi = chr(ord(ref[0]) + 1) + ref[1:]
        xs[keyingi].fill = fill_bolim

    def ajrat(*refs):
        for ref in refs:
            xs[ref].fill = fill_jami

    yozuv("A4", "Oy")
    xs["B4"] = oy_nomi
    xs["B4"].font = Font(name=FONT, size=10, color=KOK, bold=True)
    xs["B4"].alignment = Alignment(horizontal="right")
    yozuv("A5", "Yaxlitlash qadami (so'm)")
    xs["B5"] = qadam
    xs["B5"].font = f_kirit
    xs["B5"].number_format = "#,##0"
    xs["B5"].comment = Comment("Oylik shu summagacha pastga yaxlitlanadi. Asl jadvalda 1000 so'm "
                               "edi. 1 yozsangiz — yaxlitlanmaydi.", "Baliqxona",
                               width=260, height=80)
    tekshiruv(xs, "B5", type="whole", operator="greaterThanOrEqual", formula1="1",
              errorTitle="Noto'g'ri qiymat",
              error="Yaxlitlash qadami 1 yoki undan katta butun son bo'lishi kerak.")

    bolim("A7", "PUL HISOBI")
    pul = [
        ("Ajratilgan fond (guruhlar)", f"=SUM({ust(G, G_FOND)})"),
        ("Asosiy oylik (davomat bo'yicha)", f"=SUM({ust(I, I_ASOSIY)})"),
        ("+ Mukofotlar", f"=SUM({ust(I, I_MUKOFOT)})"),
        ("− Jarimalar", f"=SUM({ust(I, I_JARIMA)})"),
        ("± O'tgan oydan qoldiq", f"=SUM({ust(I, I_OTGAN)})"),
        ("Jami hisoblangan", f"=SUM({ust(I, I_JAMI)})"),
        ("− Avans berilgan", f"=SUM({ust(I, I_AVANS)})"),
        ("Qo'lga beriladi", f"=SUM({ust(I, I_QOLGA)})"),
        ("− To'landi", f"=SUM({ust(I, I_TOLANDI)})"),
        ("To'lanishi kerak (qoldiq)", f"=SUMIF({ust(I, I_QOLDIQ)},\">0\")"),
        ("Ishchilar qarzi (keyingi oyga)", f"=-SUMIF({ust(I, I_QOLDIQ)},\"<0\")"),
    ]
    for i, (matn, formula) in enumerate(pul):
        r = 8 + i
        qalin = matn in ("Jami hisoblangan", "Qo'lga beriladi", "To'lanishi kerak (qoldiq)")
        yozuv(f"A{r}", matn, qalin)
        qiymat(f"B{r}", formula, qalin=qalin)
        if qalin:
            ajrat(f"A{r}", f"B{r}")

    bolim("D7", "TEJALDI (fonddan ortgan pul)")
    tejam = [
        ("Bo'sh o'rinlar hisobiga", f"=SUM({ust(G, G_BOSH)})"),
        ("Kelmagan kunlar hisobiga", f"=SUM({ust(G, G_KELMAGAN)})"),
        ("Yaxlitlash qoldig'i", f"=SUM({ust(G, G_YAXLIT)})"),
        ("Jami tejaldi", f"=SUM({ust(G, G_TEJAM)})"),
    ]
    for i, (matn, formula) in enumerate(tejam):
        r = 8 + i
        qalin = matn == "Jami tejaldi"
        yozuv(f"D{r}", matn, qalin)
        qiymat(f"E{r}", formula, qalin=qalin)
        if qalin:
            ajrat(f"D{r}", f"E{r}")

    bolim("D13", "TEKSHIRUV")
    yozuv("D14", "Ishchilar soni (hozir / me'yor)")
    qiymat("E14", f"=COUNTIF({ust(I, I_ISM)},\"?*\")&\" / \"&SUM({ust(G, G_MEYOR)})", fmt="@")
    yozuv("D15", "Xato qatorlar («Oylik hisob»)")
    qiymat("E15", f"=COUNTIF({ust(I, I_HOLAT)},\"⚠*\")", fmt=SON)
    yozuv("D16", "Xato yozuvlar («Amallar»)")
    qiymat("E16", f"=COUNTIF({ust(A, A_TEK)},\"⚠*\")", fmt=SON)
    yozuv("D17", "Xato guruhlar («Guruhlar»)")
    qiymat("E17", f"=COUNTIF({ust(G, G_HOLAT)},\"⚠*\")", fmt=SON)
    yozuv("D18", "Fond = asosiy oylik + tejam")
    qiymat("E18", "=IF(ABS(B8-B9-E11)<0.5,\"✓ To'g'ri\",\"⚠ Farq: \"&TEXT(B8-B9-E11,\"#,##0\"))",
           fmt="@")
    yozuv("D19", "Amallar oylikka to'liq tushgan")
    # To'g'ri yozilgan amallar yig'indisi = mukofot + jarima + avans + to'lov ustunlari yig'indisi
    qiymat("E19", f"=IF(ABS(SUMIFS({ust(A, A_SUMMA)},{ust(A, A_TEK)},\"✓\")-(B10+B11+B14+B16))"
                  f"<0.5,\"✓ To'g'ri\",\"⚠ Tekshiring\")", fmt="@")
    for ref in ("E15", "E16", "E17"):
        shart_rangi(xs, ref, f"{ref}>0", "FFC7CE", "9C0006")
    for ref in ("E18", "E19"):
        shart_rangi(xs, ref, f'LEFT({ref},1)="⚠"', "FFC7CE", "9C0006")
        shart_rangi(xs, ref, f'LEFT({ref},1)="✓"', "C6EFCE", "006100")


def qollanma_varagi(qs):
    qs.column_dimensions["A"].width = 3
    qs.column_dimensions["B"].width = 110
    matnlar = [
        ("h", "Baliqxona oylik hisob-kitobi — qo'llanma"),
        ("s", "Varaqlar"),
        ("t", "Xulosa — oy nomi, yaxlitlash qadami, umumiy natijalar va avtomatik tekshiruvlar."),
        ("t", "Oylik hisob — har bir ishchining oyligi. Chop etib, imzo uchun qaydnoma sifatida "
              "ishlatish mumkin."),
        ("t", "Amallar — avans, mukofot, jarima va to'lovlar jurnali (har biri alohida qator)."),
        ("t", "Guruhlar — guruhlarning asosiy oyligi, me'yoriy soni, ish kunlari va tejam tahlili."),
        ("s", "Ranglar"),
        ("t", "Ko'k shrift — qo'lda kiritiladigan kataklar. Qora — formulalar, ularga tegmang."),
        ("t", "Qizil fon va ⚠ belgisi — xato: guruh topilmadi, ism takrorlangan yoki amal "
              "noto'g'ri yozilgan."),
        ("s", "Oylik qanday hisoblanadi"),
        ("t", "1. Guruh fondi = asosiy oylik × me'yoriy soni. Masalan, Moyka: 4 000 000 × 2 = "
              "8 000 000."),
        ("t", "2. 1 kishiga oylik = fond ÷ hozirgi odamlar soni, lekin asosiy oylikdan oshmaydi va "
              "yaxlitlash qadamigacha pastga yaxlitlanadi. Moykada 3 kishi: 8 000 000 ÷ 3 = "
              "2 666 666 → 2 666 000."),
        ("t", "3. Odam me'yordan kam bo'lsa, har kim to'liq asosiy oylikni oladi, ortgan pul "
              "«Bo'sh o'rin tejami» bo'ladi."),
        ("t", "4. Kelmagan kunlar: asosiy = 1 kishiga oylik × ishlagan kun ÷ ish kuni. Masalan, "
              "4 000 000 × 24 ÷ 26 = 3 692 307 → 3 692 000. Ushlab qolingan pul «Kelmagan kunlar "
              "tejami» bo'ladi."),
        ("t", "5. Jami hisoblangan = asosiy + mukofot − jarima ± o'tgan oydan qoldiq."),
        ("t", "6. Qo'lga beriladi = jami hisoblangan − avanslar. Qoldiq = qo'lga beriladi − "
              "to'langan."),
        ("s", "Amal yozish (misol)"),
        ("t", "Sana: 10.10.2026 | Ism: Ali | Turi: Avans | Summa: 1 000 000 | Izoh: naqd"),
        ("t", "Ism va Turi ro'yxatdan tanlanadi. Summa har doim musbat yoziladi — jarima va "
              "avans o'zi ayiriladi."),
        ("t", "Oylik kuni pul berganda «To'lov» turidagi yozuv qo'shing. Qoldiq 0 bo'lsa, holat "
              "«✓ To'landi» bo'ladi."),
        ("s", "Yangi ishchi yoki guruh qo'shish"),
        ("t", "Ishchi: «Oylik hisob» jadvalining ostidagi birinchi bo'sh qatorga ism yozing — jadval "
              "o'zi kengayadi va formulalar tushadi. Guruhni ro'yxatdan tanlang."),
        ("t", "Ishdan ketgan ishchi: qatorini o'ng tugma bilan bosing → «O'chirish» → «Jadval "
              "qatorlari». Oy o'rtasida ketgan bo'lsa, oy oxirigacha qoldiring va kelmagan "
              "kunlarini yozing."),
        ("t", "Guruh: «Guruhlar» jadvalining ostidagi bo'sh qatorga nom, oylik, me'yoriy son va ish "
              "kunini yozing. Guruh nomini o'zgartirsangiz, ishchilardagi nomni ham o'zgartiring — "
              "aks holda ⚠ ogohlantirish chiqadi."),
        ("s", "Yangi oyga o'tish"),
        ("t", "1. Faylni yangi nom bilan saqlang (masalan, Baliqxona_oylik_2026-11.xlsx)."),
        ("t", "2. «Oylik hisob»dagi «Qoldiq» ustunini nusxalab, «O'tgan oydan» ustuniga "
              "«Faqat qiymatlar» (Paste Values) sifatida qo'ying."),
        ("t", "3. «Kelmagan kun» ustunini 0 bilan to'ldiring va «Amallar» jadvalidagi yozuvlarni "
              "o'chiring."),
        ("t", "4. «Xulosa» varag'ida oy nomini o'zgartiring."),
        ("s", "Tekshiruvlar"),
        ("t", "«Xulosa» varag'idagi TEKSHIRUV bo'limida hammasi ✓ yoki 0 bo'lishi kerak. Fond "
              "tengligi buzilsa yoki amal noto'g'ri yozilsa, u yerda ⚠ chiqadi."),
    ]
    r = 1
    for tur, matn in matnlar:
        if tur == "h":
            qs.cell(r, 1, matn).font = f_sarlavha
            r += 1
        elif tur == "s":
            r += 1
            c = qs.cell(r, 1, matn)
            c.font = f_bolim
            c.fill = fill_bolim
            qs.cell(r, 2).fill = fill_bolim
        else:
            c = qs.cell(r, 2, matn)
            c.font = f_oddiy
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1


def chop_sozlash(ws, albom):
    """Chop etish: A4, eniga bitta sahifa (bo'yiga kerakricha)."""
    ws.page_setup.orientation = "landscape" if albom else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4


def kitob_yarat(chiqish, oy_nomi=OY_NOMI, qadam=QADAM, guruhlar=GURUHLAR,
                ishchilar=ISHCHILAR, amallar=AMALLAR):
    wb = Workbook()
    xs = wb.active
    xs.title = "Xulosa"
    om = wb.create_sheet("Oylik hisob")
    am = wb.create_sheet("Amallar")
    gs = wb.create_sheet("Guruhlar")
    qs = wb.create_sheet("Qo'llanma")

    wb.defined_names["OyNomi"] = DefinedName("OyNomi", attr_text="Xulosa!$B$4")
    wb.defined_names["Qadam"] = DefinedName("Qadam", attr_text="Xulosa!$B$5")

    guruhlar_varagi(gs, guruhlar)
    oylik_varagi(om, ishchilar)
    amallar_varagi(am, amallar)
    xulosa_varagi(xs, oy_nomi, qadam)
    qollanma_varagi(qs)

    for ws, albom in ((xs, False), (om, True), (am, False), (gs, True), (qs, False)):
        chop_sozlash(ws, albom)

    # Formatlanmagan kataklar ham Arial bo'lsin
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.font is not None and c.font.name != FONT:
                    yangi = copy(c.font)
                    yangi.name = FONT
                    c.font = yangi

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(chiqish)


# ---------- Hisoblangan qiymatlarni faylga joylash ----------
#
# openpyxl formulalarni qiymatsiz yozadi. Excel ochganda o'zi hisoblaydi, lekin telefon va
# messenjerlardagi oldindan ko'rish oynalari bo'sh kataklarni ko'rsatadi. Shuning uchun fayl
# LibreOffice'da hisoblanadi va natijalar asl faylga (jadval uslublari, formulalar o'zgarmagan
# holda) <v> sifatida qo'shiladi.

_KATAK = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: [a-zA-Z]+="[^"]*")*)><f>(.*?)</f>'
                    r'(?:<v\s*/>|<v>[^<]*</v>)?</c>', re.S)


def _qiymat_xml(v):
    if v is None:
        return ' t="str"', "<v></v>"
    if isinstance(v, bool):
        return ' t="b"', f"<v>{int(v)}</v>"
    if isinstance(v, (int, float)):
        return "", f"<v>{v!r}</v>" if isinstance(v, float) else f"<v>{v}</v>"
    if isinstance(v, datetime):
        seriya = (v - datetime(1899, 12, 30)).total_seconds() / 86400
        return "", f"<v>{seriya!r}</v>"
    return ' t="str"', f"<v>{escape(str(v))}</v>"


def _varaq_yollari(zf):
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "p": "http://schemas.openxmlformats.org/package/2006/relationships"}
    kitob = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    manzil = {r.get("Id"): r.get("Target") for r in rels.findall("p:Relationship", ns)}
    natija = {}
    for sh in kitob.find("m:sheets", ns):
        t = manzil[sh.get(f"{{{ns['r']}}}id")].lstrip("/")
        natija[sh.get("name")] = t if t.startswith("xl/") else "xl/" + t
    return natija


def qiymatlarni_joylash(fayl):
    """LibreOffice bilan hisoblaydi va qiymatlarni faylga yozadi. soffice bo'lmasa — o'tkazib
    yuboradi (Excel baribir ochilganda hisoblaydi)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice topilmadi: qiymatlar joylanmadi (Excel ochganda hisoblaydi).")
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{tmp}/profil",
                        "--convert-to", "xlsx", "--outdir", tmp, fayl],
                       check=True, capture_output=True, timeout=300)
        hisob = load_workbook(os.path.join(tmp, os.path.basename(fayl)), data_only=True)
        qiymatlar = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row}
                     for ws in hisob.worksheets}
    xatolar = []
    vaqtinchalik = fayl + ".tmp"
    with zipfile.ZipFile(fayl) as zin, zipfile.ZipFile(vaqtinchalik, "w",
                                                       zipfile.ZIP_DEFLATED) as zout:
        yollar = {v: k for k, v in _varaq_yollari(zin).items()}
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in yollar:
                varaq = qiymatlar[yollar[item.filename]]

                def almashtir(m):
                    ref, attrs, formula = m.group(1), m.group(2), m.group(3)
                    v = varaq.get(ref)
                    if isinstance(v, str) and v.startswith("#"):
                        xatolar.append(f"{yollar[item.filename]}!{ref}: {v}")
                    attrs = re.sub(r' t="[^"]*"', "", attrs)
                    turi, v_xml = _qiymat_xml(v)
                    return f'<c r="{ref}"{attrs}{turi}><f>{formula}</f>{v_xml}</c>'

                data = _KATAK.sub(almashtir, data.decode("utf-8")).encode("utf-8")
            zout.writestr(item, data)
    os.replace(vaqtinchalik, fayl)
    if xatolar:
        raise SystemExit("Formula xatolari: " + ", ".join(xatolar))
    return True


if __name__ == "__main__":
    chiqish = sys.argv[1] if len(sys.argv) > 1 else "Baliqxona_oylik.xlsx"
    kitob_yarat(chiqish)
    qiymatlarni_joylash(chiqish)
    print("Tayyor:", chiqish)

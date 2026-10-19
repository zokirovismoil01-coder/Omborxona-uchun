#!/usr/bin/env python3
"""Excel shabloni va ilova (js/calc.js) bir xil hisoblashini tekshiradi.

Tasodifiy oylik holatlari yaratiladi, har biri uchun Excel kitobi tuziladi va LibreOffice'da
hisoblanadi, so'ng natijalar Node.js'dagi calc.js natijalari bilan katakma-katak solishtiriladi.

Ishlatish (LibreOffice va Node.js kerak):
    python3 oylik/tools/excel_mosligi.py [holatlar_soni]
"""

import json
import os
import random
import shutil
import subprocess
import sys
import tempfile
from datetime import date

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import excel_shablon as E  # noqa: E402

CALC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "js", "calc.js")
TURLAR = {"Avans": "avans", "Mukofot": "mukofot", "Jarima": "jarima", "To'lov": "tolov"}
ISHCHI_HOLAT = {"⚠ Guruh topilmadi": "no-group", "—": "none", "Qarzdor": "debt",
                "✓ To'landi": "paid", "Qisman to'landi": "partial", "To'lanmagan": "unpaid"}


def guruh_holati(matn):
    if matn.startswith("⚠ Ish kuni"):
        return "no-days"
    if matn == "Hech kim yo'q":
        return "empty"
    if matn.startswith("Bo'sh o'rin"):
        return "vacant"
    if matn.startswith("Me'yordan ortiq"):
        return "over"
    return "ok" if matn == "Me'yorda" else matn


def holat_yarat(rnd):
    guruhlar = []
    for i in range(rnd.randint(1, 5)):
        asosiy = rnd.randint(1, 90) * 100_000 + rnd.choice([0, 0, 0, 333, 500])
        kun = rnd.choice([26, 26, 30, 31, 15, 24]) if rnd.random() > 0.05 else 0
        guruhlar.append((f"G{i}", asosiy, rnd.randint(0, 6), kun or None))
    nomlar = [g[0] for g in guruhlar]
    ishchilar = []
    for i in range(rnd.randint(0, 14)):
        guruh = rnd.choice(nomlar) if rnd.random() > 0.05 else "Yo'q guruh"
        kelmagan = 0 if rnd.random() < 0.6 else rnd.randint(1, 31)
        otgan = 0 if rnd.random() < 0.7 else rnd.randint(-30, 30) * 50_000
        ishchilar.append((f"Ishchi {i}", guruh, kelmagan, otgan, ""))
    amallar = []
    ismlar = [x[0] for x in ishchilar] + ["Begona"]
    for _ in range(rnd.randint(0, 25)):
        amallar.append((date(2026, 10, rnd.randint(1, 31)), rnd.choice(ismlar),
                        rnd.choice(list(TURLAR)), rnd.randint(1, 60) * 50_000, ""))
    qadam = rnd.choice([1, 100, 1000, 1000, 1000, 5000])
    return guruhlar, ishchilar, amallar, qadam


def js_hisob(holatlar):
    kirish = []
    for guruhlar, ishchilar, amallar, qadam in holatlar:
        kirish.append({
            "step": qadam,
            "groups": [{"id": g[0], "name": g[0], "base": g[1], "norm": g[2], "workDays": g[3] or 0}
                       for g in guruhlar],
            "members": [{"employeeId": x[0], "name": x[0], "groupId": x[1], "absent": x[2],
                         "carry": x[3]} for x in ishchilar],
            "ops": [{"employeeId": a[1], "type": TURLAR[a[2]], "amount": a[3]} for a in amallar],
        })
    skript = ("const C=require(process.argv[1]);let s='';process.stdin.on('data',d=>s+=d);"
              "process.stdin.on('end',()=>{process.stdout.write(JSON.stringify("
              "JSON.parse(s).map(x=>C.computeMonth(x))))});")
    p = subprocess.run(["node", "-e", skript, os.path.abspath(CALC)], input=json.dumps(kirish),
                       capture_output=True, text=True, check=True)
    return json.loads(p.stdout)


def main():
    soni = int(sys.argv[1]) if len(sys.argv) > 1 else 40
    rnd = random.Random(2026)
    holatlar = [holat_yarat(rnd) for _ in range(soni)]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("LibreOffice topilmadi")
    js = js_hisob(holatlar)
    xatolar = 0
    tekshirildi = 0
    with tempfile.TemporaryDirectory() as tmp:
        fayllar = []
        for i, (guruhlar, ishchilar, amallar, qadam) in enumerate(holatlar):
            f = os.path.join(tmp, f"h{i}.xlsx")
            E.kitob_yarat(f, qadam=qadam, guruhlar=guruhlar, ishchilar=ishchilar, amallar=amallar)
            fayllar.append(f)
        chiq = os.path.join(tmp, "hisob")
        subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{tmp}/profil",
                        "--convert-to", "xlsx", "--outdir", chiq] + fayllar,
                       check=True, capture_output=True, timeout=600)
        for i, f in enumerate(fayllar):
            wb = load_workbook(os.path.join(chiq, os.path.basename(f)), data_only=True)
            natija = js[i]

            def solishtir(joy, excel, jsq):
                nonlocal xatolar, tekshirildi
                tekshirildi += 1
                if excel is None:
                    excel = 0
                if excel != jsq:
                    xatolar += 1
                    print(f"holat {i}, {joy}: Excel={excel!r} JS={jsq!r}")

            om = wb["Oylik hisob"]
            sarlavha = [c.value for c in om[5]]
            qatorlar = {r["employeeId"]: r for r in natija["rows"]}
            for row in om.iter_rows(min_row=6, values_only=True):
                if not row[1]:
                    continue
                v = dict(zip(sarlavha, row))
                r = qatorlar[v["Ism"]]
                for ust, kalit in (("Asosiy", "base"), ("Mukofot", "bonus"), ("Jarima", "fine"),
                                   ("Jami hisoblangan", "total"), ("Avans", "advance"),
                                   ("Qo'lga beriladi", "toPay"), ("To'landi", "paid"),
                                   ("Qoldiq", "balance"), ("1 kishiga oylik", "perPerson")):
                    solishtir(f"{v['Ism']}.{ust}", v[ust], r[kalit])
                solishtir(f"{v['Ism']}.Holat", ISHCHI_HOLAT.get(v["Holat"], v["Holat"]),
                          r["status"])
            gs = wb["Guruhlar"]
            sarlavha = [c.value for c in gs[5]]
            guruhlar = {g["id"]: g for g in natija["groups"]}
            for row in gs.iter_rows(min_row=6, values_only=True):
                if not row[0]:
                    continue
                v = dict(zip(sarlavha, row))
                g = guruhlar[v["Guruh"]]
                for ust, kalit in (("Ajratilgan fond", "fund"), ("Hozirgi soni", "count"),
                                   ("1 kishiga oylik", "perPerson"),
                                   ("Bo'sh o'rin tejami", "vacancy"),
                                   ("Kelmagan kunlar tejami", "absence"),
                                   ("Yaxlitlash qoldig'i", "rounding"), ("Jami tejaldi", "saved")):
                    solishtir(f"{v['Guruh']}.{ust}", v[ust], g[kalit])
                solishtir(f"{v['Guruh']}.Holat", guruh_holati(v["Holat"]), g["status"])
            xs = wb["Xulosa"]
            t = natija["totals"]
            for ref, kalit in (("B8", "fund"), ("B9", "base"), ("B13", "total"),
                               ("B15", "toPay"), ("B17", "due"), ("B18", "debt"),
                               ("E11", "saved")):
                solishtir(f"Xulosa.{ref}", xs[ref].value, t[kalit])
    print(f"{soni} ta holat, {tekshirildi} ta qiymat solishtirildi, farqlar: {xatolar}")
    sys.exit(1 if xatolar else 0)


if __name__ == "__main__":
    main()
